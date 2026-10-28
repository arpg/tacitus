"""doc_write.py xlsx: a spreadsheet from CSV rows, or a copy of a template
with rows and cells filled in, written only under the EA's own folders.
A template keeps its header, styles, widths and other sheets; new rows
look like the template rows they replace; a value starting with `=` is
never a formula. Every name here is made up."""
import json
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "preprocessor"))
import doc_write as dw  # noqa: E402

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:  # the lane check below still runs
    openpyxl = None
needs_openpyxl = pytest.mark.skipif(openpyxl is None, reason="openpyxl is not installed")

HEADER = ["Your email", "Name to hire", "Category", "SpeedType", "Comments"]
OLD = [["pi@example.edu", "Ann Old", "GRA", "13000001 100%", ""],
       ["pi@example.edu", "Bo Old", "GRA", "13000002 100%", "backstop"],
       ["pi@example.edu", "Cy Old", "GRA", "13000003 100%", ""]]


def _template(tmp_path) -> Path:
    """A request form: a bold, filled header, three styled data rows, a
    wide column, a second sheet and a merged note."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Requests"
    ws.append(HEADER)
    for row in OLD:
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFDDDDDD")
    for r in (2, 3, 4):
        for c in ws[r]:
            c.font = Font(name="Arial", size=9)
    ws.column_dimensions["E"].width = 41
    ws.merge_cells("G1:H1")
    ws["G1"] = "merged note"
    wb.create_sheet("Instructions")["A1"] = "Fill one row per hire."
    p = tmp_path / "uploads" / "form.xlsx"
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(p))
    return p


def _rows(path: Path, sheet: str | None = None) -> list[list]:
    ws = openpyxl.load_workbook(str(path))[sheet] if sheet else openpyxl.load_workbook(str(path)).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


@needs_openpyxl
def test_a_new_workbook_from_csv_has_a_bold_header_and_keeps_codes_as_text(tmp_path):
    rows = dw.read_rows("﻿Name,Count,Code,Formula\nAnn,3,007,=1+1\nBo,4.5,,\n")
    out = dw.out_path(str(tmp_path / "docs" / "p" / "tally.xlsx"), ea_dir=tmp_path, suffix=".xlsx")
    res = dw.build_xlsx(rows, out, sheet="Tally")
    assert res["ok"] and res["rows_written"] == 3 and res["sheet"] == "Tally"
    ws = openpyxl.load_workbook(str(out))["Tally"]
    assert [c.value for c in ws[1]] == ["Name", "Count", "Code", "Formula"]
    assert ws["A1"].font.b and not ws["A2"].font.b and ws.freeze_panes == "A2"
    assert ws["B2"].value == 3 and ws["B3"].value == 4.5, "plain numbers are numbers"
    assert ws["C2"].value == "007", "a leading zero is a code, kept as text"
    assert ws["D2"].value == "=1+1" and ws["D2"].data_type == "s", "never a formula"
    assert ws["C3"].value is None


@needs_openpyxl
def test_a_template_keeps_its_header_styles_and_sheets_and_replace_clears_old_rows(tmp_path):
    tpl = dw.template_path(str(_template(tmp_path)), ea_dir=tmp_path, suffix=".xlsx")
    out = dw.out_path(str(tmp_path / "docs" / "update.xlsx"), ea_dir=tmp_path, suffix=".xlsx")
    new = [["pi@example.edu", "Dee New", "GRA", "13000009 100%", "new request"],
           ["pi@example.edu", "Bo Old", "GRA", "TBD", "moves to another award"]]
    res = dw.build_xlsx(new, out, template=tpl, start_row=2, replace=True)
    assert res["first_row"] == 2 and res["last_row"] == 3 and res["cells_cleared"] == 13
    got = _rows(out, "Requests")
    assert got[0][:5] == HEADER
    assert got[1][:5] == new[0] and got[2][:5] == new[1]
    assert all(v is None for v in got[3]), "the old third row is cleared"
    wb = openpyxl.load_workbook(str(out))
    ws = wb["Requests"]
    assert ws["A1"].font.b and ws["A1"].fill.fgColor.rgb == "FFDDDDDD"
    assert ws["B2"].font.name == "Arial" and ws["B2"].font.sz == 9 and not ws["B2"].font.b
    assert ws.column_dimensions["E"].width == 41
    assert "G1:H1" in [str(m) for m in ws.merged_cells.ranges] and ws["G1"].value == "merged note"
    assert wb["Instructions"]["A1"].value == "Fill one row per hire."


@needs_openpyxl
def test_rows_append_after_the_last_row_and_look_like_it(tmp_path):
    tpl = _template(tmp_path)
    out = dw.out_path(str(tmp_path / "docs" / "more.xlsx"), ea_dir=tmp_path, suffix=".xlsx")
    res = dw.build_xlsx([["pi@example.edu", "Eve Next", "GRA", "13000004", ""]], out, template=tpl)
    assert res["first_row"] == 5
    ws = openpyxl.load_workbook(str(out)).active
    assert [c.value for c in ws[2]][:2] == ["pi@example.edu", "Ann Old"], "nothing above is touched"
    assert ws["B5"].value == "Eve Next" and ws["B5"].font.name == "Arial"
    assert ws["D5"].value == "13000004", "the column holds text, so the number stays text"


@needs_openpyxl
def test_a_header_only_template_never_lends_its_header_style(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(HEADER)
    for c in wb.active[1]:
        c.font = Font(bold=True)
    tpl = tmp_path / "uploads" / "blank.xlsx"
    tpl.parent.mkdir(parents=True)
    wb.save(str(tpl))
    out = dw.out_path(str(tmp_path / "docs" / "first.xlsx"), ea_dir=tmp_path, suffix=".xlsx")
    dw.build_xlsx([["pi@example.edu", "Fay First"]], out, template=tpl)
    ws = openpyxl.load_workbook(str(out)).active
    assert ws["B2"].value == "Fay First" and not ws["B2"].font.b


@needs_openpyxl
def test_set_writes_one_cell_and_refuses_a_merged_one(tmp_path):
    tpl = _template(tmp_path)
    out = dw.out_path(str(tmp_path / "docs" / "cells.xlsx"), ea_dir=tmp_path, suffix=".xlsx")
    res = dw.build_xlsx([], out, template=tpl, cells=dw.parse_set(["e3=no backstop now", "D4=13000005"]))
    assert res["cells_set"] == 2 and res["rows_written"] == 0
    ws = openpyxl.load_workbook(str(out)).active
    assert ws["E3"].value == "no backstop now" and ws["D4"].value == "13000005"
    with pytest.raises(dw.DocError, match="merged"):
        dw.build_xlsx([], out, template=tpl, cells=dw.parse_set(["H1=x"]))
    with pytest.raises(dw.DocError, match="give a cell"):
        dw.parse_set(["Bob"])
    with pytest.raises(dw.DocError, match="no sheet 'Budget'"):
        dw.build_xlsx([], out, template=tpl, sheet="Budget")


@needs_openpyxl
def test_the_paths_stay_under_the_ea_folders(tmp_path):
    with pytest.raises(dw.DocError, match="outside"):
        dw.out_path(str(tmp_path / "elsewhere" / "a.xlsx"), ea_dir=tmp_path, suffix=".xlsx")
    with pytest.raises(dw.DocError, match="end in .xlsx"):
        dw.out_path(str(tmp_path / "docs" / "a.xlsm"), ea_dir=tmp_path, suffix=".xlsx")
    macro = tmp_path / "uploads" / "m.xlsm"
    macro.parent.mkdir(parents=True)
    macro.write_bytes(b"PK")
    with pytest.raises(dw.DocError, match="a template is a .xlsx"):
        dw.template_path(str(macro), ea_dir=tmp_path, suffix=".xlsx")


@needs_openpyxl
def test_the_cli_fills_a_template_and_reports_json(tmp_path):
    tpl = _template(tmp_path)
    (tmp_path / "state").mkdir()
    src = tmp_path / "state" / "rows.csv"
    src.write_text('pi@example.edu,Gus New,GRA,"13000006 50%, 13000007 50%",\n')
    env = {"EA_DIR": str(tmp_path), "HOME": str(Path.home()), "PATH": "/usr/bin:/bin"}
    cmd = [sys.executable, str(ROOT / "preprocessor" / "doc_write.py"), "xlsx", "--template", str(tpl),
           "--in", str(src), "--out", str(tmp_path / "docs" / "out.xlsx")]
    r = subprocess.run(cmd + ["--start-row", "2", "--replace", "--set", "E2=split"],
                       capture_output=True, text=True, env=env)
    assert r.returncode == 0, r.stdout + r.stderr
    d = json.loads(r.stdout)
    assert d["ok"] and d["rows_written"] == 1 and d["cells_set"] == 1
    assert _rows(Path(d["path"]))[1][:5] == ["pi@example.edu", "Gus New", "GRA",
                                             "13000006 50%, 13000007 50%", "split"]
    r = subprocess.run(cmd + ["--replace"], capture_output=True, text=True, env=env)
    assert r.returncode == 1 and "--replace needs --start-row" in json.loads(r.stdout)["error"]


def test_the_lanes_are_told_about_the_sheet_verbs():
    bot = (ROOT / "ea-chat-bot.py").read_text()
    assert "doc_write.py xlsx" in bot and ".xlsx" in bot.split("doc_text.py <path>)")[1][:200]
    core = (ROOT / "CLAUDE.md").read_text()
    assert "doc_write.py xlsx" in core
