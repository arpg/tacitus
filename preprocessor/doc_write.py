#!/usr/bin/env python3
"""Write a .docx from markdown, or an .xlsx from rows, under the EA's own
folders.

The chat lane allows only named scripts, so a Word file needs one: the
EA writes the content as markdown (the Write tool), then

  python3 doc_write.py docx --in <file.md> --out ~/.ea/docs/<project>/<name>.docx
      [--title "..."] [--template <existing .docx>] [--font Calibri --size 11]

`--text "..."` takes the content inline. The output must land under
~/.ea/{docs,letters,state,research}; the parent folder is made. A
template's body is replaced and its styles, headers and footers kept,
which is how a letterhead comes along.

Markdown understood: `#`..`####` headings, paragraphs (blank-line
separated), `-`/`*` bullets and `1.` numbers (two-space indent nests one
level), `>` quotes, pipe tables (the `|---|` row is the header rule),
`**bold**`, `*italic*`, `_italic_`, `` `code` ``, `---` as a page break.
Anything else is a paragraph.

A spreadsheet: the EA writes the rows as CSV (the Write tool), then

  python3 doc_write.py xlsx --in <rows.csv> --out ~/.ea/docs/<project>/<name>.xlsx
      [--template <existing .xlsx>] [--sheet <name>] [--start-row N] [--replace]
      [--set C4="text" ...]

Without a template the first CSV row is a bold header. With one, the
workbook is copied with its styles, widths and other sheets; the rows go
in at --start-row (default: after the last row with a value), each row
styled like the template row it lands on, or the nearest row above with
a value (never row 1, the header); --replace first clears the values
from --start-row down.
--set writes one cell by address, after the rows; `doc_text.py` prints a
sheet with its addresses. A value is text unless it is a plain number
in a column the template does not hold as text, and a value starting
with `=` stays text, never a formula. Output is one JSON line.
"""
from __future__ import annotations

import argparse
import bisect
import csv
import io
import json
import os
import re
import sys
from copy import copy
from pathlib import Path

EA_DIR = Path(os.environ.get("EA_DIR", str(Path.home() / ".ea")))
WRITE_ROOTS = ("docs", "letters", "state", "research")
MAX_CHARS = 2_000_000
MAX_CELLS = 200_000


class DocError(Exception):
    pass


# ── the path ──────────────────────────────────────────────────────────

def out_path(raw: str, ea_dir: Path | None = None, suffix: str = ".docx") -> Path:
    """The resolved output file: a `suffix` file under one of WRITE_ROOTS,
    parent made; DocError otherwise."""
    ea = (ea_dir or EA_DIR).expanduser()
    p = Path(raw).expanduser()
    if p.suffix.lower() != suffix:
        raise DocError(f"{raw}: the output must end in {suffix}")
    parent = p.parent.resolve()
    roots = [(ea / name).resolve() for name in WRITE_ROOTS]
    if not any(parent == r or parent.is_relative_to(r) for r in roots):
        raise DocError(f"{raw} is outside the folders doc_write writes "
                       f"(~/.ea/{{{','.join(WRITE_ROOTS)}}})")
    parent.mkdir(parents=True, exist_ok=True)
    return parent / p.name


def template_path(raw: str, ea_dir: Path | None = None, suffix: str = ".docx") -> Path:
    ea = (ea_dir or EA_DIR).expanduser()
    try:
        p = Path(raw).expanduser().resolve(strict=True)
    except (OSError, RuntimeError):
        raise DocError(f"no such template: {raw}") from None
    roots = [(ea / name).resolve() for name in WRITE_ROOTS + ("uploads",)]
    if not any(p.is_relative_to(r) for r in roots) or p.suffix.lower() != suffix:
        raise DocError(f"{raw}: a template is a {suffix} under ~/.ea")
    return p


# ── markdown ──────────────────────────────────────────────────────────

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")


def inline_runs(text: str) -> list[tuple[str, dict]]:
    """[(text, {"bold"|"italic"|"code": True})...] for one line of markdown."""
    out = []
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            out.append((piece[2:-2], {"bold": True}))
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            out.append((piece[1:-1], {"code": True}))
        elif len(piece) > 2 and piece[0] == piece[-1] and piece[0] in "*_":
            out.append((piece[1:-1], {"italic": True}))
        else:
            out.append((piece, {}))
    return out


def blocks(text: str) -> list[dict]:
    """The document as blocks: {"kind": heading|paragraph|bullet|number|
    quote|table|break, ...}. A list item keeps its "level"; a table keeps
    its "rows" (the first row is the header)."""
    out: list[dict] = []
    para: list[str] = []
    table: list[list[str]] = []

    def flush_para():
        if para:
            out.append({"kind": "paragraph", "text": " ".join(s.strip() for s in para)})
            para.clear()

    def flush_table():
        if table:
            out.append({"kind": "table", "rows": [r[:] for r in table]})
            table.clear()

    for raw in text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|") and len(stripped) > 1:
            flush_para()
            cells = [c.strip() for c in stripped[1:-1].split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c or "---") for c in cells):
                continue
            table.append(cells)
            continue
        flush_table()
        if not stripped:
            flush_para()
            continue
        m = re.match(r"(#{1,4})\s+(.*)", stripped)
        if m:
            flush_para()
            out.append({"kind": "heading", "level": len(m.group(1)), "text": m.group(2).strip()})
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            flush_para()
            out.append({"kind": "break"})
            continue
        indent = len(line) - len(line.lstrip(" "))
        m = re.match(r"[-*+]\s+(.*)", stripped)
        if m:
            flush_para()
            out.append({"kind": "bullet", "level": min(indent // 2, 2), "text": m.group(1)})
            continue
        m = re.match(r"\d+[.)]\s+(.*)", stripped)
        if m:
            flush_para()
            out.append({"kind": "number", "level": min(indent // 2, 2), "text": m.group(1)})
            continue
        if stripped.startswith(">"):
            flush_para()
            out.append({"kind": "quote", "text": stripped.lstrip("> ").strip()})
            continue
        para.append(stripped)
    flush_para()
    flush_table()
    return out


# ── the document ──────────────────────────────────────────────────────

def _style(doc, name: str, fallback: str = "Normal"):
    try:
        doc.styles[name]
        return name
    except KeyError:
        return fallback


def _add_runs(paragraph, text: str, font: str | None, size: float | None):
    from docx.shared import Pt
    for piece, fmt in inline_runs(text):
        run = paragraph.add_run(piece)
        if fmt.get("bold"):
            run.bold = True
        if fmt.get("italic"):
            run.italic = True
        if fmt.get("code"):
            run.font.name = "Consolas"
        elif font:
            run.font.name = font
        if size and not fmt.get("code"):
            run.font.size = Pt(size)


def _clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def build(text: str, out: Path, *, title: str = "", template: Path | None = None,
          font: str | None = None, size: float | None = None) -> dict:
    from docx import Document
    if len(text) > MAX_CHARS:
        raise DocError(f"the content is over {MAX_CHARS:,} characters")
    doc = Document(str(template)) if template else Document()
    if template:
        _clear_body(doc)
    if font or size:
        from docx.shared import Pt
        normal = doc.styles["Normal"]
        if font:
            normal.font.name = font
        if size:
            normal.font.size = Pt(size)
    counts = {"headings": 0, "paragraphs": 0, "lists": 0, "tables": 0}
    if title:
        doc.add_heading(title, level=0)
        counts["headings"] += 1
    for b in blocks(text):
        kind = b["kind"]
        if kind == "heading":
            p = doc.add_heading("", level=b["level"])
            _add_runs(p, b["text"], None, None)
            counts["headings"] += 1
        elif kind == "paragraph":
            p = doc.add_paragraph()
            _add_runs(p, b["text"], font, size)
            counts["paragraphs"] += 1
        elif kind in ("bullet", "number"):
            base = "List Bullet" if kind == "bullet" else "List Number"
            name = base + (f" {b['level'] + 1}" if b["level"] else "")
            p = doc.add_paragraph(style=_style(doc, name, _style(doc, base, "List Paragraph")))
            _add_runs(p, b["text"], font, size)
            counts["lists"] += 1
        elif kind == "quote":
            p = doc.add_paragraph(style=_style(doc, "Quote", "Normal"))
            _add_runs(p, b["text"], font, size)
            if p.style.name == "Normal":
                for r in p.runs:
                    r.italic = True
            counts["paragraphs"] += 1
        elif kind == "table":
            rows = b["rows"]
            width = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=width)
            t.style = _style(doc, "Table Grid", "Normal Table")
            for i, row in enumerate(rows):
                for j in range(width):
                    cell = t.cell(i, j)
                    cell.text = ""
                    _add_runs(cell.paragraphs[0], row[j] if j < len(row) else "", font, size)
                    if i == 0:
                        for r in cell.paragraphs[0].runs:
                            r.bold = True
            counts["tables"] += 1
        elif kind == "break":
            doc.add_page_break()
    doc.save(str(out))
    return {"ok": True, "path": str(out), **counts}


# ── the spreadsheet ──────────────────────────────────────────────────

_NUMBER = re.compile(r"-?(0|[1-9][0-9]{0,14})(\.[0-9]+)?")


def read_rows(text: str, tab: bool = False) -> list[list[str]]:
    """The CSV (or TSV) as rows of strings; a row of empty fields stays a
    blank row."""
    rows = list(csv.reader(io.StringIO(text.lstrip("\ufeff")), delimiter="\t" if tab else ","))
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def parse_set(items: list[str]) -> list[tuple[str, str]]:
    """[("C4", "text")...] from `--set C4=text` arguments."""
    from openpyxl.utils.cell import coordinate_from_string
    out = []
    for item in items or []:
        ref, sep, value = item.partition("=")
        ref = ref.strip().upper()
        try:
            if not sep:
                raise ValueError
            coordinate_from_string(ref)
        except ValueError:
            raise DocError(f"--set {item!r}: give a cell and its value, like C4=Jane Doe") from None
        out.append((ref, value))
    return out


def _put(cell, raw: str, keep_text: bool):
    """Write one value: a plain number as a number unless the column holds
    text, everything else (a leading `=` too) as text."""
    if raw == "":
        cell.value = None
    elif not keep_text and _NUMBER.fullmatch(raw.strip()):
        n = raw.strip()
        cell.value = float(n) if "." in n else int(n)
    else:
        cell.value = raw
        cell.data_type = "s"


def _has_value(ws, r: int) -> bool:
    return any(c.value not in (None, "") for c in ws[r])


def _look(ws, r: int, width: int) -> list[tuple]:
    """Each column's style in row r, and whether the column holds text."""
    out = []
    for c in range(1, width + 1):
        cell = ws.cell(r, c)
        out.append((copy(cell._style), cell.data_type == "s" and cell.value not in (None, "")))
    return out


def build_xlsx(rows: list[list[str]], out: Path, *, template: Path | None = None,
               sheet: str = "", start_row: int | None = None, replace: bool = False,
               cells: list[tuple[str, str]] | None = None, header: bool = True) -> dict:
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    if sum(len(r) for r in rows) + len(cells or []) > MAX_CELLS:
        raise DocError(f"the content is over {MAX_CELLS:,} cells")
    if template:
        wb = openpyxl.load_workbook(template)
        if sheet and sheet not in wb.sheetnames:
            raise DocError(f"no sheet {sheet!r} in the template; it has: {', '.join(wb.sheetnames)}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = (sheet or "Sheet1")[:31]
    filled = [r for r in range(1, ws.max_row + 1) if _has_value(ws, r)] if template else []
    first = start_row or (filled[-1] + 1 if filled else 1)
    if first < 1:
        raise DocError("--start-row counts from 1")
    width = max([len(r) for r in rows] + [ws.max_column if template else 0])

    # Each new row is styled like the template row it lands on, or the
    # nearest row above with a value; never like row 1, the header. Read
    # before --replace clears anything.
    looks: dict[int, list[tuple]] = {}
    source: dict[int, int] = {}
    for r in range(first, first + len(rows)):
        k = bisect.bisect_right(filled, r)
        src = filled[k - 1] if k else 0
        if src > 1:
            source[r] = src
            if src not in looks:
                looks[src] = _look(ws, src, width)

    cleared = 0
    if replace and template:
        for r in range(first, ws.max_row + 1):
            for c in ws[r]:
                if isinstance(c, MergedCell):
                    continue
                if c.value not in (None, ""):
                    cleared += 1
                c.value = None
    for i, row in enumerate(rows):
        r = first + i
        src = source.get(r, 0)
        look = looks.get(src, [])
        if src and src != r and ws.row_dimensions[src].height:
            ws.row_dimensions[r].height = ws.row_dimensions[src].height
        for j, raw in enumerate(row, start=1):
            cell = ws.cell(r, j)
            if isinstance(cell, MergedCell):
                raise DocError(f"{cell.coordinate} is inside a merged range; write its top-left cell")
            style, text_col = look[j - 1] if j <= len(look) else (None, False)
            if style is not None and src != r:
                cell._style = copy(style)
            _put(cell, raw, text_col)
    if not template:
        if rows and header:
            for cell in ws[first]:
                cell.font = Font(bold=True)
            ws.freeze_panes = ws.cell(first + 1, 1).coordinate
        for j in range(1, width + 1):
            longest = max((len(r[j - 1]) for r in rows if j <= len(r)), default=0)
            ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 2, 8), 60)
    for ref, value in cells or []:
        cell = ws[ref]
        if isinstance(cell, MergedCell):
            raise DocError(f"{ref} is inside a merged range; write its top-left cell")
        _put(cell, value, cell.data_type == "s" and cell.value not in (None, ""))
    wb.save(str(out))
    return {"ok": True, "path": str(out), "sheet": ws.title,
            "rows_written": len(rows), "first_row": first if rows else None,
            "last_row": first + len(rows) - 1 if rows else None,
            "cells_cleared": cleared, "cells_set": len(cells or [])}


# ── CLI ───────────────────────────────────────────────────────────────

def cmd_docx(args) -> int:
    if bool(args.inp) == bool(args.text):
        print(json.dumps({"ok": False, "error": "give --in <file> or --text, not both"}))
        return 2
    try:
        if args.inp:
            src = Path(args.inp).expanduser()
            if not src.is_file():
                raise DocError(f"no such input: {args.inp}")
            text = src.read_text(encoding="utf-8", errors="replace")
        else:
            text = args.text
        out = out_path(args.out)
        tpl = template_path(args.template) if args.template else None
        res = build(text, out, title=args.title or "", template=tpl, font=args.font, size=args.size)
    except DocError as e:
        print(json.dumps({"ok": False, "error": str(e)}))
        return 1
    except Exception as e:  # python-docx errors: a bad template, a locked file
        print(json.dumps({"ok": False, "error": f"{type(e).__name__}: {str(e)[:200]}"}))
        return 1
    print(json.dumps(res))
    return 0


def cmd_xlsx(args) -> int:
    if not args.inp and not args.set:
        print(json.dumps({"ok": False, "error": "give --in <rows.csv>, --set CELL=value, or both"}))
        return 2
    try:
        rows: list[list[str]] = []
        if args.inp:
            src = Path(args.inp).expanduser()
            if not src.is_file():
                raise DocError(f"no such input: {args.inp}")
            if src.stat().st_size > MAX_CHARS:
                raise DocError(f"the input is over {MAX_CHARS:,} bytes")
            rows = read_rows(src.read_text(encoding="utf-8", errors="replace"),
                             tab=src.suffix.lower() == ".tsv")
        cells = parse_set(args.set)
        if args.replace and not args.start_row:
            raise DocError("--replace needs --start-row: the first row to clear and write")
        out = out_path(args.out, suffix=".xlsx")
        tpl = template_path(args.template, suffix=".xlsx") if args.template else None
        res = build_xlsx(rows, out, template=tpl, sheet=args.sheet or "", start_row=args.start_row,
                         replace=args.replace, cells=cells, header=not args.no_header)
    except DocError as e:
        print(json.dumps({"ok": False, "error": str(e)}))
        return 1
    except Exception as e:  # openpyxl errors: a bad template, a locked file
        print(json.dumps({"ok": False, "error": f"{type(e).__name__}: {str(e)[:200]}"}))
        return 1
    print(json.dumps(res))
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0], allow_abbrev=False)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("docx", help="write a .docx from markdown or text", allow_abbrev=False)
    p.add_argument("--in", dest="inp", help="the markdown or text file")
    p.add_argument("--text", help="the content inline")
    p.add_argument("--out", required=True, help="the .docx to write, under ~/.ea/{docs,letters,state,research}")
    p.add_argument("--title", default="", help="a title heading above the content")
    p.add_argument("--template", help="an existing .docx whose styles, headers and footers to keep")
    p.add_argument("--font", help="the body font, e.g. Calibri")
    p.add_argument("--size", type=float, help="the body size in points")
    p.set_defaults(func=cmd_docx)
    p = sub.add_parser("xlsx", help="write an .xlsx from CSV rows, or fill a template", allow_abbrev=False)
    p.add_argument("--in", dest="inp", help="the rows as CSV (.tsv: tab-separated)")
    p.add_argument("--out", required=True, help="the .xlsx to write, under ~/.ea/{docs,letters,state,research}")
    p.add_argument("--template", help="an existing .xlsx to copy, styles and other sheets kept")
    p.add_argument("--sheet", help="the sheet to write (default: the first; names a new workbook's sheet)")
    p.add_argument("--start-row", type=int, help="the row the first CSV row goes to "
                   "(default: after the last row with a value)")
    p.add_argument("--replace", action="store_true", help="clear the values from --start-row down first")
    p.add_argument("--set", action="append", metavar="CELL=VALUE", help="write one cell, e.g. C4=Jane Doe; repeatable")
    p.add_argument("--no-header", action="store_true", help="a new workbook: the first row is not a bold header")
    p.set_defaults(func=cmd_xlsx)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
